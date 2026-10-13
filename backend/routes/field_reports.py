import io
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from services.auth import get_current_user, is_admin_user, is_mastec_byers_admin
from services.geocode import geocode

router = APIRouter(tags=["field-reports"])
_db = None


def set_db(db):
    global _db
    _db = db


# Column layout (order matters — matches the reference sheet and the UI table)
DFR_COLUMNS = [
    ("date", "Date", 11), ("job", "Job", 12), ("market", "Market", 9), ("sow", "SOW", 9), ("job_type", "Job Type", 10),
    ("job_address", "Job Address", 30), ("fielder", "Fielder at Site", 18), ("travel_hrs", "Travel to site (hrs)", 10),
    ("work_hrs", "Work @job (hrs)", 10), ("return_hrs", "Return to Home/next Site (hrs)", 12), ("miles", "Total miles driven", 10),
    ("footage", "Job footage (ft)", 11), ("poles", "# Poles", 8), ("computer_hrs", "Time on computer for Job", 12),
    ("status", "Status", 14), ("comments", "Comments", 40),
]


class FieldReportIn(BaseModel):
    date: str
    fielder: Optional[str] = ""
    rows: List[Dict[str, Any]] = Field(default_factory=list)
    track: List[Dict[str, Any]] = Field(default_factory=list)  # [{lat, lng, t, row_id}]


def _is_admin(user: dict) -> bool:
    return is_admin_user(user) or is_mastec_byers_admin(user)


async def _get_report(report_id: str, user: dict, write: bool = False) -> dict:
    """Any company user may view a report; only its owner or an admin may change it."""
    doc = await _db.field_reports.find_one({"id": report_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Report not found")
    if write and doc.get("user_id") != user["id"] and not _is_admin(user):
        raise HTTPException(status_code=403, detail="Only the report owner or an admin can change it")
    return doc


def _rows_date(rows: List[dict]) -> Optional[str]:
    """The one date every filled-in row carries, or None when rows disagree / are blank."""
    dates = {str(r.get("date") or "").strip() for r in (rows or []) if any(str(r.get(k) or "").strip() for k in ("job", "job_id", "job_address", "status", "comments", "miles", "footage", "poles"))}
    dates.discard("")
    return next(iter(dates)) if len(dates) == 1 else None


async def _effective_date(user_id: str, report_id: Optional[str], date: str, rows: List[dict]) -> str:
    """A report belongs to the date its rows show. Falls back to `date` when another report already owns that day."""
    rd = _rows_date(rows)
    if not rd or rd == date:
        return date
    q: Dict[str, Any] = {"user_id": user_id, "date": rd}
    if report_id:
        q["id"] = {"$ne": report_id}
    return date if await _db.field_reports.find_one(q, {"_id": 0, "id": 1}) else rd


async def migrate_report_dates(db) -> int:
    """One-time fix: old reports whose rows all show a different date move to that date."""
    n = 0
    async for doc in db.field_reports.find({}, {"_id": 0, "id": 1, "user_id": 1, "date": 1, "rows": 1}):
        eff = await _effective_date(doc.get("user_id"), doc["id"], doc.get("date"), doc.get("rows") or [])
        if eff != doc.get("date"):
            await db.field_reports.update_one({"id": doc["id"]}, {"$set": {"date": eff}})
            n += 1
    return n


@router.get("/field-reports")
async def list_field_reports(date: Optional[str] = None, all: bool = False, current_user: dict = Depends(get_current_user)):
    """Own reports by default; ?all=1 lists every company user's reports."""
    q: Dict[str, Any] = {}
    if not all:
        q["user_id"] = current_user["id"]
    if date:
        q["date"] = date
    docs = await _db.field_reports.find(q, {"_id": 0, "track": 0}).sort("date", -1).to_list(500)
    return docs


@router.post("/field-reports", status_code=201)
async def create_field_report(body: FieldReportIn, current_user: dict = Depends(get_current_user)):
    existing = await _db.field_reports.find_one({"user_id": current_user["id"], "date": body.date}, {"_id": 0})
    if existing:
        return existing
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid.uuid4()), "user_id": current_user["id"], "user_email": current_user.get("email", ""),
        "date": body.date, "fielder": body.fielder or "", "rows": body.rows, "track": body.track,
        "created_at": now, "updated_at": now,
    }
    await _db.field_reports.insert_one(dict(doc))
    return doc


class HomeIn(BaseModel):
    address: str = ""
    lat: Optional[float] = None
    lng: Optional[float] = None


@router.get("/field-reports/home")
async def get_home(current_user: dict = Depends(get_current_user)):
    """Fielder's saved home / start address (per user)."""
    doc = await _db.users.find_one({"id": current_user["id"]}, {"_id": 0, "dfr_home": 1})
    return (doc or {}).get("dfr_home") or {"address": "", "lat": None, "lng": None}


@router.put("/field-reports/home")
async def set_home(body: HomeIn, current_user: dict = Depends(get_current_user)):
    home = {"address": (body.address or "").strip(), "lat": body.lat, "lng": body.lng}
    await _db.users.update_one({"id": current_user["id"]}, {"$set": {"dfr_home": home}})
    return home


@router.get("/field-reports/geocode")
async def geocode_address(q: str, current_user: dict = Depends(get_current_user)):
    q = (q or "").strip()
    if len(q) < 4:
        raise HTTPException(status_code=400, detail="Enter a fuller address")
    hit = await geocode(q)
    if not hit:
        raise HTTPException(status_code=404, detail="Address not found — try adding city and state")
    return hit


def _require_admin(user: dict):
    """Team Reports: admins, or users an admin has flagged as supervisor."""
    if not (is_admin_user(user) or is_mastec_byers_admin(user) or user.get("is_supervisor")):
        raise HTTPException(status_code=403, detail="Admin or supervisor access required")


async def _team_rows(date_from: str, date_to: str):
    """All fielders' DFR rows between two dates (inclusive), each row tagged with fielder / user email / date."""
    if date_to < date_from:
        date_from, date_to = date_to, date_from
    docs = await _db.field_reports.find({"date": {"$gte": date_from, "$lte": date_to}}, {"_id": 0, "track": 0}).sort([("date", 1), ("fielder", 1)]).to_list(5000)
    users = {u["id"]: u for u in await _db.users.find({}, {"_id": 0, "id": 1, "email": 1, "name": 1, "full_name": 1}).to_list(5000)}
    rows: List[Dict[str, Any]] = []
    for d in docs:
        u = users.get(d.get("user_id"), {})
        who = d.get("fielder") or u.get("full_name") or u.get("name") or u.get("email") or ""
        for r in d.get("rows") or []:
            if not any(str(r.get(k) or "").strip() for k in ("job", "job_id", "job_address", "work_hrs", "travel_hrs", "miles", "comments")):
                continue
            rows.append({**r, "date": r.get("date") or d.get("date"), "fielder": r.get("fielder") or who, "report_fielder": who, "user_email": u.get("email", ""), "report_id": d.get("id")})
    return rows, date_from, date_to


@router.get("/field-reports/team")
async def team_field_reports(date_from: str, date_to: str, current_user: dict = Depends(get_current_user)):
    _require_admin(current_user)
    rows, f, t = await _team_rows(date_from, date_to)
    totals = {k: round(sum(_num(r.get(k)) for r in rows), 2) for k in TOTAL_KEYS}
    fielders = sorted({r["report_fielder"] for r in rows if r.get("report_fielder")})
    return {"date_from": f, "date_to": t, "rows": rows, "totals": totals, "fielders": fielders, "count": len(rows)}


@router.get("/field-reports/team/export")
async def team_field_reports_export(date_from: str, date_to: str, current_user: dict = Depends(get_current_user)):
    _require_admin(current_user)
    rows, f, t = await _team_rows(date_from, date_to)
    if not rows:
        raise HTTPException(status_code=404, detail="No report rows between those dates")
    wb = build_dfr_workbook({"rows": rows, "date": f, "fielder": "Team"}, title=f"Team {f} to {t}")
    return _xlsx_response(wb, f"Team_Field_Report_{f}_to_{t}.xlsx")


@router.get("/field-reports/{report_id}")
async def get_field_report(report_id: str, current_user: dict = Depends(get_current_user)):
    return await _get_report(report_id, current_user)


@router.put("/field-reports/{report_id}")
async def update_field_report(report_id: str, body: FieldReportIn, current_user: dict = Depends(get_current_user)):
    doc = await _get_report(report_id, current_user, write=True)
    date = await _effective_date(doc.get("user_id"), report_id, body.date, body.rows)
    upd = {"date": date, "fielder": body.fielder or "", "rows": body.rows, "track": body.track,
           "updated_at": datetime.now(timezone.utc).isoformat(), "updated_by": current_user.get("email", "")}
    await _db.field_reports.update_one({"id": report_id}, {"$set": upd})
    return await _db.field_reports.find_one({"id": report_id}, {"_id": 0})


@router.delete("/field-reports/{report_id}")
async def delete_field_report(report_id: str, current_user: dict = Depends(get_current_user)):
    await _get_report(report_id, current_user, write=True)
    await _db.field_reports.delete_one({"id": report_id})
    return {"deleted": True}


def _fmt(v):
    if v is None or v == "":
        return ""
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v


TOTAL_KEYS = ("travel_hrs", "work_hrs", "return_hrs", "miles", "footage", "poles", "computer_hrs")


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def build_dfr_workbook(report: dict, title: str = "Daily Field Report") -> Workbook:
    """One sheet: header, one line per used row (no padding), then a TOTALS line."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    header_fill = PatternFill("solid", fgColor="000000")
    grey_fill = PatternFill("solid", fgColor="D9D9D9")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="000000")
    thick = Side(style="medium", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for ci, (_, label, width) in enumerate(DFR_COLUMNS, 1):
        c = ws.cell(row=1, column=ci, value=label)
        c.fill = header_fill
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 78
    rows = [r for r in (report.get("rows") or []) if any(str(r.get(k) or "").strip() for k, _, _ in DFR_COLUMNS if k != "date")]
    for ri, row in enumerate(rows):
        for ci, (key, _, _) in enumerate(DFR_COLUMNS, 1):
            val = row.get(key, "") if key != "date" else (row.get("date") or report.get("date") or "")
            c = ws.cell(row=ri + 2, column=ci, value=_fmt(val))
            c.border = border
            c.alignment = Alignment(horizontal="center" if key not in ("job_address", "comments", "status") else "left", vertical="center", wrap_text=True)
            if ri % 2 == 0:
                c.fill = grey_fill
    # TOTALS line
    tr = len(rows) + 2
    totals = {k: sum(_num(r.get(k)) for r in rows) for k in TOTAL_KEYS}
    for ci, (key, _, _) in enumerate(DFR_COLUMNS, 1):
        val = "TOTALS" if key == "date" else (_fmt(round(totals[key], 2)) if key in totals else "")
        c = ws.cell(row=tr, column=ci, value=val)
        c.fill = total_fill
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(left=thin, right=thin, top=thick, bottom=thick)
    n = len(rows) + 1  # data rows + totals
    # thick separators after Fielder and after Return (as on the reference sheet)
    for ci in (7, 10):
        for ri in range(1, n + 2):
            cell = ws.cell(row=ri, column=ci)
            cell.border = Border(left=cell.border.left, right=thick, top=cell.border.top, bottom=cell.border.bottom)
    ws.freeze_panes = "A2"
    return wb


def _xlsx_response(wb: Workbook, fname: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/field-reports/export/range")
async def export_field_reports_range(date_from: str, date_to: str, current_user: dict = Depends(get_current_user)):
    """One combined sheet for all of the caller's reports between two dates (inclusive), rows in date order, with totals."""
    if date_to < date_from:
        date_from, date_to = date_to, date_from
    docs = await _db.field_reports.find({"user_id": current_user["id"], "date": {"$gte": date_from, "$lte": date_to}},
                                        {"_id": 0, "track": 0}).sort("date", 1).to_list(1000)
    rows: List[Dict[str, Any]] = []
    fielders: List[str] = []
    for d in docs:
        for r in d.get("rows") or []:
            rows.append({**r, "date": r.get("date") or d.get("date")})
        if d.get("fielder") and d["fielder"] not in fielders:
            fielders.append(d["fielder"])
    if not rows:
        raise HTTPException(status_code=404, detail="No report rows between those dates")
    wb = build_dfr_workbook({"rows": rows, "date": date_from, "fielder": ", ".join(fielders)}, title=f"DFR {date_from} to {date_to}")
    fielder = (fielders[0] if len(fielders) == 1 else "combined").replace(" ", "_")
    return _xlsx_response(wb, f"Daily_Field_Report_{date_from}_to_{date_to}_{fielder}.xlsx")


@router.get("/field-reports/{report_id}/export/xlsx")
async def export_field_report_xlsx(report_id: str, current_user: dict = Depends(get_current_user)):
    report = await _get_report(report_id, current_user)
    wb = build_dfr_workbook(report)
    fielder = (report.get("fielder") or "").replace(" ", "_") or "fielder"
    return _xlsx_response(wb, f"Daily_Field_Report_{report.get('date', '')}_{fielder}.xlsx")


# ---------------- Cost report (fielder pay) ----------------
DEFAULT_COST_RATES = {"job_fee": 75.0, "per_pole": 6.50, "per_mile": 0.76, "per_foot": 0.06, "free_feet": 1000}


async def get_cost_rates() -> Dict[str, float]:
    """Admin-editable pay rates (settings doc), falling back to the defaults."""
    doc = await _db.settings.find_one({"id": "cost_rates"}, {"_id": 0}) or {}
    rates = dict(DEFAULT_COST_RATES)
    for k in rates:
        if isinstance(doc.get(k), (int, float)):
            rates[k] = float(doc[k])
    return rates


class CostRatesIn(BaseModel):
    job_fee: float = Field(ge=0)
    per_pole: float = Field(ge=0)
    per_mile: float = Field(ge=0)
    per_foot: float = Field(ge=0)
    free_feet: float = Field(ge=0)


@router.get("/settings/cost-rates")
async def read_cost_rates(current_user: dict = Depends(get_current_user)):
    return {"rates": await get_cost_rates(), "defaults": DEFAULT_COST_RATES}


@router.put("/settings/cost-rates")
async def write_cost_rates(body: CostRatesIn, current_user: dict = Depends(get_current_user)):
    if not _is_admin(current_user):
        raise HTTPException(status_code=403, detail="Admin access required")
    rates = {k: float(v) for k, v in body.model_dump().items()}
    await _db.settings.update_one({"id": "cost_rates"}, {"$set": {**rates, "id": "cost_rates", "updated_by": current_user.get("email"),
                                                                   "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    return {"rates": rates, "defaults": DEFAULT_COST_RATES}


def _num(v) -> float:
    try:
        return float(str(v).replace(",", "").strip() or 0)
    except ValueError:
        return 0.0


HOUR_FIELDS = ("travel_hrs", "work_hrs", "return_hrs", "computer_hrs")


def _row_hours(r: dict) -> float:
    """Total hours on a job row: travel + work @ job + return + computer time."""
    return round(sum(_num(r.get(k)) for k in HOUR_FIELDS), 2)


def _row_filled(r: dict) -> bool:
    return any(str(r.get(k) or "").strip() for k in ("job", "job_id", "job_address", "miles", "footage", "poles"))


def cost_line(line: dict, rates: Dict[str, float], hourly_rate: float = 0.0) -> dict:
    """Pay for one job row: job fee (if toggled on) + per-pole + per-mile + per-foot over the free footage + hours x hourly rate."""
    poles, miles, feet = _num(line.get("poles")), _num(line.get("miles")), _num(line.get("footage"))
    hours = _num(line.get("hours"))
    excess = max(0.0, feet - rates["free_feet"])
    job_fee = rates["job_fee"] if line.get("job_fee", True) else 0.0
    pole_pay = round(poles * rates["per_pole"], 2)
    mile_pay = round(miles * rates["per_mile"], 2)
    foot_pay = round(excess * rates["per_foot"], 2)
    hours_pay = round(hours * hourly_rate, 2)
    return {**line, "poles": poles, "miles": miles, "footage": feet, "hours": hours, "excess_feet": excess, "job_fee_amount": job_fee,
            "pole_pay": pole_pay, "mile_pay": mile_pay, "foot_pay": foot_pay, "hours_pay": hours_pay,
            "total": round(job_fee + pole_pay + mile_pay + foot_pay + hours_pay, 2)}


@router.get("/field-reports/cost-report/rows")
async def cost_report_rows(date_from: str, date_to: str, current_user: dict = Depends(get_current_user)):
    """Every filled job row on the caller's reports between two dates (inclusive) — the pick list for the cost report."""
    if date_to < date_from:
        date_from, date_to = date_to, date_from
    docs = await _db.field_reports.find({"user_id": current_user["id"], "date": {"$gte": date_from, "$lte": date_to}},
                                        {"_id": 0, "track": 0}).sort("date", 1).to_list(1000)
    out: List[Dict[str, Any]] = []
    for d in docs:
        for i, r in enumerate(d.get("rows") or []):
            if not _row_filled(r):
                continue
            out.append({"key": f"{d['id']}:{r.get('id') or i}", "date": r.get("date") or d.get("date"), "job": r.get("job") or "",
                        "job_id": r.get("job_id") or "", "job_address": r.get("job_address") or "", "poles": _num(r.get("poles")),
                        "miles": _num(r.get("miles")), "footage": _num(r.get("footage")), "hours": _row_hours(r), "fielder": d.get("fielder") or ""})
    return {"rows": out, "rates": await get_cost_rates()}


class CostReportIn(BaseModel):
    date_from: str
    date_to: str
    lines: List[Dict[str, Any]] = Field(default_factory=list)  # selected rows, each with job_fee: bool
    hourly_rate: float = Field(default=0.0, ge=0)  # entered by the user each time the report is generated


def build_cost_workbook(date_from: str, date_to: str, lines: List[dict], fielder: str, rates: Dict[str, float], hourly_rate: float = 0.0) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cost Report"
    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="1F2937")
    thin = Side(style="thin", color="D4D4D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '"$"#,##0.00'
    ws["A1"] = f"Cost Report  {date_from} to {date_to}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Fielder: {fielder}" if fielder else ""
    r = rates
    ws["A3"] = f"Rates: ${r['job_fee']:.2f} per job · ${r['per_pole']:.2f} per pole · ${r['per_mile']:.2f} per mile · ${r['per_foot']:.2f} per ft over {int(r['free_feet']):,} ft · ${hourly_rate:.2f} per hour"
    ws["A3"].font = Font(italic=True, color="52525B")
    cols = [("date", "Date", 11), ("job", "Job", 14), ("job_address", "Job Address", 32), ("job_fee_amount", "Job Fee", 10),
            ("poles", "# Poles", 8), ("pole_pay", "Pole Pay", 10), ("miles", "Miles", 8), ("mile_pay", "Mileage Pay", 11),
            ("footage", "Footage (ft)", 11), ("excess_feet", f"Ft over {int(rates['free_feet']):,}", 11), ("foot_pay", "Footage Pay", 11),
            ("hours", "Hours", 8), ("hours_pay", "Hours Pay", 11), ("total", "Line Total", 12)]
    head = 5
    for c, (_, title, width) in enumerate(cols, start=1):
        cell = ws.cell(row=head, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = width
    money_keys = {"job_fee_amount", "pole_pay", "mile_pay", "foot_pay", "hours_pay", "total"}
    for i, line in enumerate(lines, start=head + 1):
        for c, (key, _, _) in enumerate(cols, start=1):
            cell = ws.cell(row=i, column=c, value=line.get(key, ""))
            cell.border = border
            if key in money_keys:
                cell.number_format = money
    tr = head + len(lines) + 1
    ws.cell(row=tr, column=1, value="TOTAL").font = bold
    for c, (key, _, _) in enumerate(cols, start=1):
        if key in money_keys or key in ("poles", "miles", "footage", "excess_feet", "hours"):
            cell = ws.cell(row=tr, column=c, value=round(sum(_num(l.get(key)) for l in lines), 2))
            cell.font = bold
            if key in money_keys:
                cell.number_format = money
    # Summary: one line per pay component, then the overall total
    sr = tr + 2
    ws.cell(row=sr, column=1, value="SUMMARY").font = bold
    summary = [("Job fees", "job_fee_amount"), ("Pole pay", "pole_pay"), ("Mileage pay", "mile_pay"), ("Footage pay", "foot_pay"),
               (f"Hours pay ({round(sum(_num(l.get('hours')) for l in lines), 2):g} hrs x ${hourly_rate:.2f})", "hours_pay")]
    for i, (label, key) in enumerate(summary, start=sr + 1):
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=3, value=round(sum(_num(l.get(key)) for l in lines), 2))
        cell.number_format = money
    ws.cell(row=sr + len(summary) + 1, column=1, value="OVERALL TOTAL").font = bold
    cell = ws.cell(row=sr + len(summary) + 1, column=3, value=round(sum(_num(l.get("total")) for l in lines), 2))
    cell.font = bold
    cell.number_format = money
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    return wb


@router.post("/field-reports/cost-report")
async def cost_report(body: CostReportIn, format: str = "json", current_user: dict = Depends(get_current_user)):
    """Line-item pay breakdown for the selected job rows; ?format=xlsx downloads it."""
    rates = await get_cost_rates()
    lines = [cost_line(l, rates, body.hourly_rate) for l in body.lines]
    lines.sort(key=lambda l: (str(l.get("date") or ""), str(l.get("job") or "")))
    total = round(sum(l["total"] for l in lines), 2)
    if format == "xlsx":
        fielder = next((l.get("fielder") for l in lines if l.get("fielder")), current_user.get("name") or "")
        wb = build_cost_workbook(body.date_from, body.date_to, lines, fielder, rates, body.hourly_rate)
        return _xlsx_response(wb, f"Cost_Report_{body.date_from}_to_{body.date_to}.xlsx")
    return {"lines": lines, "total": total, "rates": rates, "hourly_rate": body.hourly_rate,
            "hours": round(sum(l["hours"] for l in lines), 2),
            "subtotals": {k: round(sum(l[k] for l in lines), 2) for k in ("job_fee_amount", "pole_pay", "mile_pay", "foot_pay", "hours_pay")}}
