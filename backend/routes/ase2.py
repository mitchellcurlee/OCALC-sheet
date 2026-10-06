import io
import os
from datetime import datetime, timezone
from typing import Any, Dict, List

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from services.auth import get_current_user

router = APIRouter(tags=["ase2"])
_db = None
TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "templates", "ase2_template.xlsm")

YN = ["Yes", "No"]
# Fielding Checklist sheet — exactly the questions / cells / dropdowns of the uploaded template (column B holds the answer)
SECTIONS: List[Dict[str, Any]] = [
    {"title": "SITE VISIT INFO", "fields": [
        {"key": "business_name", "cell": "B2", "label": "Name of Business?"},
        {"key": "site_address", "cell": "B4", "label": "Site Address? Incl Suite/Bldg #"},
        {"key": "lcon_name", "cell": "B6", "label": "LCON Name?"},
        {"key": "lcon_email", "cell": "B8", "label": "LCON Email?"},
        {"key": "lcon_phone", "cell": "B10", "label": "LCON Phone Number?"},
        {"key": "local_contact_name", "cell": "B12", "label": "Local Contact Name?"},
        {"key": "local_contact_number", "cell": "B14", "label": "Local Contact Number?"},
        {"key": "availability", "cell": "B16", "label": "Site Visit Availability Dates/Times?"},
    ]},
    {"title": "CUSTOMER & SITE VISIT CONTACT INFORMATION", "fields": [
        {"key": "fielder", "cell": "B19", "label": "Fielder Who Completed Site Visit?"},
        {"key": "visit_date", "cell": "B21", "label": "Date of Site Visit?", "type": "date"},
        {"key": "business_name_2", "cell": "B23", "label": "Name of Business?"},
        {"key": "owner_contact", "cell": "B25", "label": "Owner or Leasing Contact Name and Number?"},
        {"key": "site_address_2", "cell": "B27", "label": "Site Address? Incl Suite/Bldg #"},
        {"key": "contact_met", "cell": "B29", "label": "Contact Met at Location?"},
        {"key": "contact_phone", "cell": "B31", "label": "Contact Phone Number?"},
        {"key": "contact_position", "cell": "B33", "label": "Contact Position?"},
    ]},
    {"title": "BUILDING ATTRIBUTES", "fields": [
        {"key": "building_type", "cell": "B36", "label": "Building Type?", "options": ["Retail", "Industrial", "Office", "Sports", "Government", "Land", "Mulit-Family", "Central Office", "Health Care", "Hospitality", "Strip Mall", "Specialty"]},
        {"key": "multi_tenant", "cell": "B38", "label": "Multi-Tenant Building?", "options": YN},
        {"key": "tenant_count", "cell": "B39", "label": "If yes, number of tenants?"},
        {"key": "floors", "cell": "B41", "label": "Number of Floors in Building? incl basement/attic"},
        {"key": "telco_rooms", "cell": "B43", "label": "Number of Building Telco Rooms?", "options": ["Single Telco Room", "Multiple Telco Rooms"]},
    ]},
    {"title": "TELCO ROOM ATTRIBUTES", "fields": [
        {"key": "path_type", "cell": "B46", "label": "Path Type from ROW to Telco Room?", "options": ["Aerial", "Buried"]},
        {"key": "telco_location", "cell": "B48", "label": "Customer Telco Room Location?", "options": ["North", "South", "East", "West"]},
        {"key": "telco_floor", "cell": "B50", "label": "Customer Telco Room Floor?"},
        {"key": "ceiling_height", "cell": "B52", "label": "Customer Telco Room Ceiling Height?", "options": ["10ft or Less", "Greater than 10ft"]},
        {"key": "environment", "cell": "B54", "label": "Telco Room Environment?", "options": ["INSIDE (Temp Controlled)", "INSIDE (No Temp Control)", "OUTSIDE (Non-Insulated)"]},
    ]},
    {"title": "EXISTING EQUIPMENT ATTRIBUTES", "fields": [
        {"key": "equip_location", "cell": "B57", "label": "Existing Equipment Location?", "options": ["Outside Wall Mount", "H-Frame", "PED (Specify Location)", "Rack/Shelf", "Inside Wall Mount", "Other (Specify)"]},
        {"key": "weatherhead", "cell": "B58", "label": "If outside, weatherhead?", "options": ["N/A", "Existing", "Customer Place New", "AT&T Place New"]},
        {"key": "att_equipment", "cell": "B61", "label": "Existing AT&T Equipment Present?", "options": YN},
        {"key": "att_ids", "cell": "B62", "label": "If yes, IDs/CLLIs?"},
        {"key": "bidi_retrofit", "cell": "B63", "label": "BIDI Retrofit?", "options": ["YES", "NO"]},
        {"key": "firm_commitment", "cell": "B64", "label": "Customer firm commitment (Date and time, 10 + days out)"},
        {"key": "commitment_contact", "cell": "B65", "label": "Contact name and number"},
        {"key": "mounting", "cell": "B66", "label": "Existing Equipment Mounting?", "options": ["Wall Mount (Backboard)", "Rack/Shelf Mount"]},
        {"key": "wall_location", "cell": "B69", "label": "If Wall Mount (Backboard) — Existing Equipment Mounting Location?", "options": ["North", "South", "East", "West"]},
        {"key": "room_for_new", "cell": "B70", "label": "Is There Room to Place New Equipment? 2FTx2FT", "options": YN},
        {"key": "rack_location", "cell": "B73", "label": "If Rack/Shelf Mount — Existing Equipment Mounting Location?", "options": ["North", "South", "East", "West"]},
        {"key": "rack_bay", "cell": "B74", "label": "Rack Bay?"},
        {"key": "rack_number", "cell": "B75", "label": "Rack Number?"},
        {"key": "power_source", "cell": "B77", "label": "Existing Power Source?", "options": ["AC (Wall)", "AC (Rack)", "DC (Rack)", "Twist Lock", "No Existing Power"]},
        {"key": "outlet_type", "cell": "B78", "label": "If yes, power outlet type?", "options": ["Standard 3 Prong", "110V L5 R15 Amp Twist Lock Wall Mount", "220V L6 R30 Amp Twist Lock Rack Mount"]},
        {"key": "ground_present", "cell": "B80", "label": "Existing Ground Present?", "options": YN},
        {"key": "ground_source", "cell": "B81", "label": "If yes, ground source?", "options": ["Electrical Panel", "Ground Rod", "Building Steel", "Unk Need To Verify", "Not Needed (Plastic FBT)"]},
        {"key": "conduit_available", "cell": "B83", "label": "Existing Conduit Available?", "options": ["Yes (Spare Duct)", "No (None AT&T Can Use)"]},
        {"key": "conduit_size", "cell": "B84", "label": "If yes, conduit size?", "options": ["2in", "4in", "6in"]},
        {"key": "same_floor", "cell": "B86", "label": "Proposed Equipment Location on Same Floor as Customer Router?", "options": YN},
        {"key": "router_floor", "cell": "B87", "label": "If no, what floor is cutomer router on?"},
        {"key": "within_300", "cell": "B89", "label": "Proposed Equipment Location within 300ft of Customer Router?", "options": YN},
    ]},
    {"title": "SITE REQUIREMENTS & MEASUREMENTS", "fields": [
        {"key": "room_ready", "cell": "B92", "label": "Room Ready*", "options": YN, "note": "Backboard space, power, tagged ground source, & conduit must all be present to show Room Ready"},
        {"key": "dist_existing_to_proposed", "cell": "B95", "label": "Distance from Existing Equipment to Proposed Equipment Location?", "note": "ADD a minimum of 25 feet when FTE is on wall & NTE is in rack"},
        {"key": "dist_poe_to_proposed", "cell": "B98", "label": "Distance from Point Of Entry to Proposed Equipment Location?", "note": "ADD a minimum of 25 feet when path goes up & over wall/door/rack/etc."},
        {"key": "dist_proposed_to_outlet", "cell": "B101", "label": "Distance from  Proposed Equipment Location to Existing Outlet?", "note": "If an existing power outlet is further than 5 feet from proposed NTE, then a new outlet is required"},
    ]},
    {"title": "ADDITIONAL NOTES", "fields": [
        {"key": "job_notes", "cell": "B105", "label": "JOB NOTES*", "type": "textarea"},
        {"key": "misc_notes", "cell": "B119", "label": "MISC. Fielding Notes", "type": "textarea"},
    ]},
]
REQ_ITEMS = [("backboard", "Backboard"), ("wall_outlet", "Wall Outlet"), ("pullbox", "Pullbox"), ("conduit", "Conduit (ROW to MPOE)"), ("weatherhead_eyebolt", "Weatherhead and Eyebolt")]
REQ_GROUPS = [("cust_req", "Customer Requirements*", 107, "For Non-EF jobs, *OR* jobs that need service extended to the server room with prior AT&T approval ONLY"),
              ("att_req", "AT&T Requirements*", 113, "For EF jobs *AND* for ANY new equipment being placed. \"AT&T place power & ground, FTE, Ciena 3903x & DEMARC\" etc.")]
# Handoff Site Details sheet — one "Label: value" line per row (A3..A29)
HANDOFF = [("site_id", "Site ID", 3), ("common_id", "Common ID", 4), ("asr_number", "ASR Number", 5), ("pon", "PON", 6), ("crdd", "CRDD", 7),
           ("site_type", "Site Type", 9), ("process_type", "Process Type", 10), ("business_unit", "Business Unit", 11), ("e_rate", "E-Rate", 12),
           ("customer", "Customer", 13), ("end_customer", "End Customer", 14), ("acna", "ACNA", 15),
           ("address", "Address", 17), ("sub_address", "Sub Address", 18), ("city", "City", 19), ("state", "State", 20), ("zip", "ZIP", 21),
           ("lata", "LATA", 22), ("customer_clli", "Customer CLLI", 23), ("swc_clli", "SWC CLLI", 24),
           ("nss_pm", "NSS PM", 26), ("bupm", "BUPM", 27), ("lcon", "LCON", 28), ("secondary_lcon", "Secondary LCON", 29)]


# Customer Call Script & Response sheet — verbatim text of the template
CALL_SCRIPT = {
    "reach_title": "If you reach someone:",
    "reach": "Hi, my name is Brooke Robinson and I am with Byers Engineering Company, a subcontractor of AT&T. I am calling in regards to a ASE fiber order I recieved for address ____________________ and you are listed as a contact for this order. Are you available to meet with an inspector, ____________________________? The inspector will need access to the MPOE (main point of entry) for the telecommunication equipment into the building and the place where the equipment is or will be placed.",
    "voicemail_title": "When Leaving a Voicemail:",
    "voicemail": "Hi, my name is Brooke Robinson and I am with Byers Engineering Company, a subcontractor of AT&T. I am calling in regards to a ASE fiber order I recieved for address ____________________ and would like to know if you are available to meet with an inspector on (give date). Again this is Brooke Robinson and I am calling in regards to a ASE fiber order I recieved for ________. I can be reached at 256-499-8223 or you can respond to the email sent previously. Thank you",
    "typical_title": "Typical Responses to Questions",
}


def set_db(db):
    global _db
    _db = db


@router.get("/ase2-schema")
async def ase2_schema(current_user: dict = Depends(get_current_user)):
    return {"sections": SECTIONS, "req_items": [{"key": k, "label": l} for k, l in REQ_ITEMS], "req_groups": [{"key": k, "label": l, "note": n} for k, l, _, n in REQ_GROUPS],
            "handoff": [{"key": k, "label": l} for k, l, _ in HANDOFF], "call_script": CALL_SCRIPT}


@router.get("/jobs/{job_id}/ase2-data")
async def get_ase2(job_id: str, current_user: dict = Depends(get_current_user)):
    from routes.jobs import can_view_job
    job = await can_view_job(job_id, current_user)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    return job.get("ase2_data") or {}


class Ase2Body(BaseModel):
    data: Dict[str, Any]


@router.put("/jobs/{job_id}/ase2-data")
async def save_ase2(job_id: str, body: Ase2Body, current_user: dict = Depends(get_current_user)):
    from routes.jobs import can_view_job
    if not await can_view_job(job_id, current_user):
        raise HTTPException(status_code=404, detail="Job not found")
    await _db.jobs.update_one({"id": job_id}, {"$set": {"ase2_data": body.data, "updated_at": datetime.now(timezone.utc).isoformat()}})
    return body.data


def _req_line(label: str, val: str) -> str:
    v = (val or "").upper()
    return f"{label} - {v}" if v in ("YES", "NO") else f"{label} -                                             YES          NO"


@router.get("/jobs/{job_id}/export/ase2")
async def export_ase2(job_id: str, current_user: dict = Depends(get_current_user)):
    """Fill the uploaded 'ASE Site Visit - Blank Template' workbook (macros kept) with the saved ASE Form 2 answers."""
    import openpyxl
    from routes.jobs import can_view_job
    job = await can_view_job(job_id, current_user)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    d = job.get("ase2_data") or {}
    wb = openpyxl.load_workbook(TEMPLATE, keep_vba=True)
    ws = wb["Fielding Checklist"]
    for sec in SECTIONS:
        for f in sec["fields"]:
            ws[f["cell"]] = d.get(f["key"], "") or ""
    for gkey, _, start, _ in REQ_GROUPS:
        for i, (ikey, ilabel) in enumerate(REQ_ITEMS):
            ws[f"B{start + i}"] = _req_line(ilabel, d.get(f"{gkey}_{ikey}", ""))
    ho = wb["Handoff Site Details"]
    for key, label, row in HANDOFF:
        ho[f"A{row}"] = f"{label}: {d.get('handoff_' + key, '') or ''}"
    if d.get("customer_response"):
        wb["Customer Call Script & Response"]["B16"] = d["customer_response"]
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in (job.get("site_id") or job.get("name") or "Job"))
    return StreamingResponse(buf, media_type="application/vnd.ms-excel.sheet.macroEnabled.12",
                             headers={"Content-Disposition": f'attachment; filename="ASE_Site_Visit_{safe}.xlsm"'})


class LabelSizeBody(BaseModel):
    size: str


@router.put("/jobs/{job_id}/label-size")
async def set_all_label_sizes(job_id: str, body: LabelSizeBody, current_user: dict = Depends(get_current_user)):
    """Bulk smart-label size for every pole + underground item on the job (normal | small | off)."""
    from routes.jobs import can_view_job
    if body.size not in ("normal", "small", "off"):
        raise HTTPException(status_code=400, detail="size must be normal, small or off")
    if not await can_view_job(job_id, current_user):
        raise HTTPException(status_code=404, detail="Job not found")
    r1 = await _db.poles.update_many({"job_id": job_id}, {"$set": {"label_size": body.size}})
    r2 = await _db.underground.update_many({"job_id": job_id}, {"$set": {"label_size": body.size}})
    return {"poles": r1.modified_count, "underground": r2.modified_count, "size": body.size}
