# Excel bulk import: template download + job import (Poles + Underground sheets)
import io
from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from models.pole import Pole
from models.underground import Underground

router = APIRouter()

_db = None

def set_db(database):
    global _db
    _db = database


POLE_HEADERS = ["Pole #", "Latitude", "Longitude", "Size", "Class", "Condition", "Owner", "Notes"]
UG_HEADERS = ["Type", "Label", "Latitude", "Longitude", "Notes"]


@router.get("/import/template")
async def download_import_template():
    """Blank .xlsx template with Poles + Underground sheets"""
    wb = Workbook()
    ws_p = wb.active
    ws_p.title = "Poles"
    ws_p.append(POLE_HEADERS)
    ws_u = wb.create_sheet("Underground")
    ws_u.append(UG_HEADERS)
    header_fill = PatternFill("solid", fgColor="DC2626")
    for ws in (ws_p, ws_u):
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for col_cells in ws.columns:
            ws.column_dimensions[col_cells[0].column_letter].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="pole_import_template.xlsx"'},
    )


def _cell(row, idx):
    if idx >= len(row) or row[idx] is None:
        return None
    v = str(row[idx]).strip()
    return v or None


def _num(row, idx):
    if idx >= len(row) or row[idx] is None:
        return None
    try:
        return float(str(row[idx]).strip())
    except (TypeError, ValueError):
        return None


@router.post("/jobs/{job_id}/import-excel")
async def import_excel(job_id: str, file: UploadFile = File(...), mode: str = Form("add")):
    """Import poles + underground attributes from a filled template. mode: add | replace"""
    job = await _db.jobs.find_one({"id": job_id})
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if mode not in ("add", "replace"):
        raise HTTPException(status_code=400, detail="mode must be 'add' or 'replace'")
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read file — upload a valid .xlsx spreadsheet")

    sheet_names = {n.strip().lower(): n for n in wb.sheetnames}
    ws_p = wb[sheet_names["poles"]] if "poles" in sheet_names else None
    ws_u = wb[sheet_names["underground"]] if "underground" in sheet_names else None
    if ws_p is None and ws_u is None:
        raise HTTPException(status_code=400, detail="Spreadsheet needs a 'Poles' and/or 'Underground' sheet (use the template)")

    if mode == "replace":
        await _db.poles.delete_many({"job_id": job_id})
        await _db.underground.delete_many({"job_id": job_id})
        await _db.connections.delete_many({"job_id": job_id})
        next_order = 0
    else:
        orders = [p.get("sort_order") or 0 for p in await _db.poles.find({"job_id": job_id}, {"sort_order": 1}).to_list(5000)]
        orders += [u.get("sort_order") or 0 for u in await _db.underground.find({"job_id": job_id}, {"sort_order": 1}).to_list(5000)]
        next_order = (max(orders) + 1) if orders else 0

    errors = []
    pole_docs, ug_docs = [], []

    if ws_p is not None:
        for i, row in enumerate(ws_p.iter_rows(min_row=2, values_only=True), start=2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            pole_number = _cell(row, 0)
            if not pole_number:
                errors.append(f"Poles row {i}: missing Pole # — skipped")
                continue
            pole = Pole(
                job_id=job_id,
                pole_number=pole_number,
                latitude=_num(row, 1),
                longitude=_num(row, 2),
                pole_size=_cell(row, 3),
                pole_class=_cell(row, 4),
                condition=_cell(row, 5),
                pole_owner=_cell(row, 6),
                notes=_cell(row, 7),
                sort_order=next_order,
            )
            next_order += 1
            pole_docs.append(pole.model_dump())

    if ws_u is not None:
        for i, row in enumerate(ws_u.iter_rows(min_row=2, values_only=True), start=2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            ug_type = _cell(row, 0)
            if not ug_type:
                errors.append(f"Underground row {i}: missing Type — skipped")
                continue
            ug = Underground(
                job_id=job_id,
                type=ug_type,
                label=_cell(row, 1),
                latitude=_num(row, 2),
                longitude=_num(row, 3),
                notes=_cell(row, 4),
                sort_order=next_order,
            )
            next_order += 1
            ug_docs.append(ug.model_dump())

    if pole_docs:
        await _db.poles.insert_many(pole_docs)
    if ug_docs:
        await _db.underground.insert_many(ug_docs)

    return {
        "mode": mode,
        "poles_created": len(pole_docs),
        "underground_created": len(ug_docs),
        "errors": errors,
    }
