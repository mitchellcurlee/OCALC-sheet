# Excel bulk import + End Path persistence integration tests (iteration 30)
import io
import os
import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://utility-pole-map.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "Admin@mastec.com"
ADMIN_PASSWORD = "Mastecbyers12345"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def auth_headers(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def temp_job(auth_headers):
    # create temp job for import tests
    payload = {"name": "TEST_ITER30_IMPORT", "job_number": "TEST-ITER30", "location": "QA"}
    r = requests.post(f"{API}/jobs", json=payload, headers=auth_headers, timeout=30)
    assert r.status_code in (200, 201), r.text
    job = r.json()
    job["_auth"] = auth_headers
    yield job
    # cleanup
    requests.delete(f"{API}/jobs/{job['id']}", headers=auth_headers, timeout=30)


# -------- Template download --------

class TestTemplate:
    def test_get_template_200_xlsx(self):
        r = requests.get(f"{API}/import/template", timeout=30)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct or "octet-stream" in ct, ct
        # validate structure
        wb = load_workbook(io.BytesIO(r.content))
        assert "Poles" in wb.sheetnames
        assert "Underground" in wb.sheetnames
        poles_headers = [c.value for c in wb["Poles"][1]]
        assert poles_headers == ["Pole #", "Latitude", "Longitude", "Size", "Class", "Condition", "Owner", "Notes"]
        ug_headers = [c.value for c in wb["Underground"][1]]
        assert ug_headers == ["Type", "Label", "Latitude", "Longitude", "Notes"]


def _build_xlsx(pole_rows, ug_rows):
    wb = Workbook()
    ws_p = wb.active
    ws_p.title = "Poles"
    ws_p.append(["Pole #", "Latitude", "Longitude", "Size", "Class", "Condition", "Owner", "Notes"])
    for row in pole_rows:
        ws_p.append(row)
    ws_u = wb.create_sheet("Underground")
    ws_u.append(["Type", "Label", "Latitude", "Longitude", "Notes"])
    for row in ug_rows:
        ws_u.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# -------- Import add mode --------

class TestImportAddMode:
    def test_import_add_mode(self, temp_job):
        job_id = temp_job["id"]
        xlsx = _build_xlsx(
            pole_rows=[
                ["TEST_P1", 33.111, -84.111, "40", "3", "Good", "GA Power", "n1"],
                ["TEST_P2", 33.112, -84.112, "40", "3", "Good", "GA Power", "n2"],
                [None, 33.113, -84.113, "40", "3", "Good", "GA Power", "bad row missing pole #"],
            ],
            ug_rows=[
                ["Handhole", "TEST_HH1", 33.114, -84.114, "note"],
            ],
        )
        files = {"file": ("import.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {"mode": "add"}
        r = requests.post(f"{API}/jobs/{job_id}/import-excel", files=files, data=data, timeout=60)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["poles_created"] == 2, body
        assert body["underground_created"] == 1, body
        assert len(body["errors"]) == 1, body
        assert "row 4" in body["errors"][0].lower() or "missing" in body["errors"][0].lower()

        # verify persistence via GET (job poles requires auth)
        poles = requests.get(f"{API}/jobs/{job_id}/poles", headers=temp_job["_auth"], timeout=30).json()
        pole_numbers = {p["pole_number"] for p in poles}
        assert {"TEST_P1", "TEST_P2"}.issubset(pole_numbers), pole_numbers

        ugs = requests.get(f"{API}/jobs/{job_id}/underground", headers=temp_job["_auth"], timeout=30).json()
        labels = {u.get("label") for u in ugs}
        assert "TEST_HH1" in labels, labels


# -------- Import replace mode (non-cumulative) --------

class TestImportReplaceMode:
    def test_import_replace_mode(self, temp_job):
        job_id = temp_job["id"]
        # ensure some existing data (from previous add test or add fresh)
        pre_poles = requests.get(f"{API}/jobs/{job_id}/poles", timeout=30).json()
        pre_count = len(pre_poles)
        # first add some junk poles that should be wiped by replace
        junk = _build_xlsx(
            pole_rows=[["JUNK_1", 30.0, -80.0, None, None, None, None, None]],
            ug_rows=[],
        )
        requests.post(f"{API}/jobs/{job_id}/import-excel",
                      files={"file": ("j.xlsx", junk, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                      data={"mode": "add"}, timeout=60)

        # Now REPLACE with exactly 2 poles + 1 UG
        xlsx = _build_xlsx(
            pole_rows=[
                ["REP_P1", 33.201, -84.201, "40", "3", "Good", "GA Power", ""],
                ["REP_P2", 33.202, -84.202, "40", "3", "Good", "GA Power", ""],
            ],
            ug_rows=[["Handhole", "REP_HH1", 33.203, -84.203, ""]],
        )
        r = requests.post(f"{API}/jobs/{job_id}/import-excel",
                          files={"file": ("r.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                          data={"mode": "replace"}, timeout=60)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["mode"] == "replace"
        assert body["poles_created"] == 2
        assert body["underground_created"] == 1

        poles = requests.get(f"{API}/jobs/{job_id}/poles", headers=temp_job["_auth"], timeout=30).json()
        pole_numbers = {p["pole_number"] for p in poles}
        # ONLY the new 2 remain — old TEST_P1/TEST_P2/JUNK_1 gone
        assert pole_numbers == {"REP_P1", "REP_P2"}, pole_numbers
        assert len(poles) == 2

        ugs = requests.get(f"{API}/jobs/{job_id}/underground", headers=temp_job["_auth"], timeout=30).json()
        assert len(ugs) == 1
        assert ugs[0].get("label") == "REP_HH1"


# -------- Import validation errors --------

class TestImportErrors:
    def test_import_invalid_mode(self, temp_job):
        job_id = temp_job["id"]
        xlsx = _build_xlsx([["X", 33.0, -84.0, None, None, None, None, None]], [])
        r = requests.post(f"{API}/jobs/{job_id}/import-excel",
                          files={"file": ("x.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                          data={"mode": "bogus"}, timeout=30)
        assert r.status_code == 400, r.text

    def test_import_not_xlsx(self, temp_job):
        job_id = temp_job["id"]
        r = requests.post(f"{API}/jobs/{job_id}/import-excel",
                          files={"file": ("hello.txt", io.BytesIO(b"not-excel"), "text/plain")},
                          data={"mode": "add"}, timeout=30)
        assert r.status_code == 400, r.text

    def test_import_job_not_found(self):
        xlsx = _build_xlsx([["X", 33.0, -84.0, None, None, None, None, None]], [])
        r = requests.post(f"{API}/jobs/does-not-exist/import-excel",
                          files={"file": ("x.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                          data={"mode": "add"}, timeout=30)
        assert r.status_code == 404, r.text


# -------- End Path persistence --------

class TestEndPathPersist:
    def test_pole_end_path_toggle(self, temp_job):
        job_id = temp_job["id"]
        # create a pole via API (POST /poles with job_id in body)
        create = requests.post(f"{API}/poles", json={
            "job_id": job_id, "pole_number": "EP_P1", "latitude": 33.5, "longitude": -84.5,
            "pole_size": "40", "condition": "Good"
        }, timeout=30)
        assert create.status_code in (200, 201), create.text
        pole = create.json()
        pid = pole["id"]

        # toggle end_path
        r = requests.put(f"{API}/poles/{pid}", json={"end_path": True}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json().get("end_path") is True

        # verify GET reflects it
        got = requests.get(f"{API}/jobs/{job_id}/poles", headers=temp_job["_auth"], timeout=30).json()
        me = next((p for p in got if p["id"] == pid), None)
        assert me and me.get("end_path") is True

        # toggle back off
        r2 = requests.put(f"{API}/poles/{pid}", json={"end_path": False}, timeout=30)
        assert r2.status_code == 200, r2.text
        assert r2.json().get("end_path") is False

    def test_underground_end_path_toggle(self, temp_job):
        job_id = temp_job["id"]
        create = requests.post(f"{API}/underground", json={
            "job_id": job_id, "type": "Handhole", "label": "EP_HH1", "latitude": 33.51, "longitude": -84.51
        }, timeout=30)
        assert create.status_code in (200, 201), create.text
        ug = create.json()
        uid = ug["id"]

        r = requests.put(f"{API}/underground/{uid}", json={"end_path": True}, timeout=30)
        assert r.status_code == 200, r.text
        assert r.json().get("end_path") is True

        got = requests.get(f"{API}/jobs/{job_id}/underground", headers=temp_job["_auth"], timeout=30).json()
        me = next((u for u in got if u["id"] == uid), None)
        assert me and me.get("end_path") is True


# -------- ASE date_of_site_visit persistence --------

class TestASEDate:
    def test_ase_date_of_site_visit_persists(self, temp_job):
        job_id = temp_job["id"]
        payload = {"ase_data": {"date_of_site_visit": "2026-01-15", "other_note": "qa"}}
        r = requests.put(f"{API}/jobs/{job_id}/ase-data", json=payload, headers=temp_job["_auth"], timeout=30)
        assert r.status_code == 200, r.text
        got = requests.get(f"{API}/jobs/{job_id}/ase-data", headers=temp_job["_auth"], timeout=30)
        assert got.status_code == 200
        ase = got.json().get("ase_data", {})
        assert ase.get("date_of_site_visit") == "2026-01-15"
