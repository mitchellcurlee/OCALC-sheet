"""Backend tests for the photo-upload / pole save / gallery / Excel-export fixes.

Verifies:
  - POST /api/upload-photo (no auth) returns {url, path}
  - GET /api/files/{path} publicly serves bytes
  - POST /api/poles with multiple photo URLs succeeds (no DocumentTooLarge)
  - GET /api/poles/{id}/gallery returns HTML with <img> per photo
  - GET /api/jobs/{job_id}/export/excel returns xlsx, contains PHOTOS row +
    feet/inches formatted heights
"""
import io
import os
import re
import uuid
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
ADMIN_EMAIL = "tim.gables@byers.com"
ADMIN_PASSWORD = "Byers@13430"
SEED_JOB_ID = "e5dc0bae-96bb-4cef-8001-7baf4e41d80a"

# 1x1 PNG
PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\xcf"
    b"\xc0\x00\x00\x00\x03\x00\x01\xa1\x06\x00>\x00\x00\x00\x00IEND\xaeB`\x82"
)


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
        timeout=30,
    )
    if r.status_code != 200:
        pytest.skip(f"Admin login failed: {r.status_code} {r.text}")
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


@pytest.fixture(scope="module")
def uploaded_urls():
    """Upload 3 small photos via the public endpoint, return list of URLs."""
    urls = []
    for i in range(3):
        files = {"file": (f"t_{i}.png", io.BytesIO(PNG_BYTES), "image/png")}
        r = requests.post(f"{BASE_URL}/api/upload-photo", files=files, timeout=60)
        assert r.status_code == 200, f"upload {i} failed: {r.status_code} {r.text}"
        data = r.json()
        assert "url" in data and data["url"].startswith("/api/files/")
        assert "path" in data
        urls.append(data["url"])
    return urls


class TestUploadAndServe:
    def test_upload_no_auth(self):
        files = {"file": ("a.png", io.BytesIO(PNG_BYTES), "image/png")}
        r = requests.post(f"{BASE_URL}/api/upload-photo", files=files, timeout=60)
        assert r.status_code == 200
        j = r.json()
        assert j["url"].startswith("/api/files/")
        # Now fetch it back
        r2 = requests.get(f"{BASE_URL}{j['url']}", timeout=30)
        assert r2.status_code == 200
        assert r2.headers.get("content-type", "").startswith("image/")
        assert len(r2.content) > 0

    def test_upload_accepts_image_returns_valid_url(self, uploaded_urls):
        # Fetch each uploaded URL and assert image bytes round-trip
        for u in uploaded_urls:
            r = requests.get(f"{BASE_URL}{u}", timeout=30)
            assert r.status_code == 200
            assert len(r.content) > 0


class TestPoleSaveWithPhotos:
    """Reproduces the original bug: saving a pole with several photo URLs."""

    pole_id = None

    def test_create_pole_with_photos(self, admin_headers, uploaded_urls):
        payload = {
            "job_id": SEED_JOB_ID,
            "pole_number": f"TEST_{uuid.uuid4().hex[:6]}",
            "pole_size": "40-3",
            "condition": "Good",
            "ocalc_left_photo": uploaded_urls[0],
            "ocalc_middle_photo": uploaded_urls[1],
            "ocalc_right_photo": uploaded_urls[2],
            "birthmark_photo": uploaded_urls[0],
            "mgnv_photo": uploaded_urls[1],
            "down_guys": [
                {"size": "3/8", "height": 35.5, "lead_length": "10", "direction": "N"},
            ],
            "knuckles": [{"photo": uploaded_urls[2]}],
            "attachments": [
                {"company": "HOA", "height": 22.25, "photo": uploaded_urls[0]},
            ],
        }
        r = requests.post(
            f"{BASE_URL}/api/poles", json=payload, headers=admin_headers, timeout=60
        )
        assert r.status_code in (200, 201), f"create failed: {r.status_code} {r.text}"
        data = r.json()
        assert "id" in data
        TestPoleSaveWithPhotos.pole_id = data["id"]

        # Verify persistence + photo URLs round-tripped
        g = requests.get(
            f"{BASE_URL}/api/poles/{data['id']}", headers=admin_headers, timeout=30
        )
        assert g.status_code == 200
        body = g.json()
        assert body["ocalc_left_photo"] == uploaded_urls[0]
        assert body["birthmark_photo"] == uploaded_urls[0]
        assert body["mgnv_photo"] == uploaded_urls[1]
        # Heights preserved as numeric
        assert float(body["down_guys"][0]["height"]) == 35.5
        assert float(body["attachments"][0]["height"]) == 22.25

    def test_pole_gallery_html(self, admin_headers):
        assert TestPoleSaveWithPhotos.pole_id, "previous test must succeed"
        r = requests.get(
            f"{BASE_URL}/api/poles/{TestPoleSaveWithPhotos.pole_id}/gallery",
            timeout=30,
        )
        assert r.status_code == 200
        assert "text/html" in r.headers.get("content-type", "")
        # Should contain several <img tags (one per photo)
        imgs = re.findall(r"<img[^>]+src=", r.text)
        assert len(imgs) >= 4, f"expected >=4 imgs, got {len(imgs)}"

    def test_excel_export_contains_photos_and_ft_in(self, admin_headers):
        r = requests.get(
            f"{BASE_URL}/api/jobs/{SEED_JOB_ID}/export/excel",
            headers=admin_headers,
            timeout=60,
        )
        assert r.status_code == 200
        ct = r.headers.get("content-type", "")
        assert "spreadsheet" in ct or "officedocument" in ct, ct

        # Open the xlsx, look for PHOTOS row + ft/in heights
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(r.content))
        all_text = []
        hyperlinks = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        all_text.append(str(cell.value))
                    if cell.hyperlink is not None:
                        hyperlinks.append(cell.hyperlink.target or "")
        joined = "\n".join(all_text)
        assert "PHOTOS" in joined, "PHOTOS row not found in Excel"
        # Some pole height should be in ft/in format like 35' 6"
        assert re.search(r"\d+'\s*\d+\"", joined), "no ft/in formatted height found"
        # gallery hyperlink for newly created pole
        if TestPoleSaveWithPhotos.pole_id:
            assert any(
                f"/api/poles/{TestPoleSaveWithPhotos.pole_id}/gallery" in h
                for h in hyperlinks
            ), "expected gallery hyperlink not present"

    def test_cleanup_pole(self, admin_headers):
        if TestPoleSaveWithPhotos.pole_id:
            requests.delete(
                f"{BASE_URL}/api/poles/{TestPoleSaveWithPhotos.pole_id}",
                headers=admin_headers,
                timeout=30,
            )
