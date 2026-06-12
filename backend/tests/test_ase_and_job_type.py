"""Tests for ASE Site Walk Survey endpoints and Job Type field."""
import os
import io
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
ADMIN_EMAIL = "admin@mastecbyers.com"
ADMIN_PASSWORD = "Mastecbyers12345"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json().get("access_token") or r.json().get("token")


@pytest.fixture(scope="module")
def auth_headers(token):
    return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}


@pytest.fixture(scope="module")
def created_job_id(auth_headers):
    """Create a test job with job_type=SVCL/FS for ASE testing."""
    payload = {
        "name": "TEST_ASE_JOB",
        "description": "Testing ASE export",
        "location": "Test Location",
        "job_type": "SVCL/FS",
    }
    r = requests.post(f"{BASE_URL}/api/jobs", json=payload, headers=auth_headers, timeout=30)
    assert r.status_code == 200, r.text
    job = r.json()
    yield job["id"]
    # Cleanup
    requests.delete(f"{BASE_URL}/api/jobs/{job['id']}", headers=auth_headers, timeout=30)


# ============= Job Type tests =============

class TestJobType:
    def test_create_job_with_svcl_fs(self, auth_headers):
        payload = {"name": "TEST_JT_SVCL", "job_type": "SVCL/FS"}
        r = requests.post(f"{BASE_URL}/api/jobs", json=payload, headers=auth_headers, timeout=30)
        assert r.status_code == 200, r.text
        job = r.json()
        job_id = job["id"]
        try:
            # Verify persisted via GET
            g = requests.get(f"{BASE_URL}/api/jobs/{job_id}", headers=auth_headers, timeout=30)
            assert g.status_code == 200
            assert g.json().get("job_type") == "SVCL/FS", f"Expected SVCL/FS, got {g.json().get('job_type')}"
        finally:
            requests.delete(f"{BASE_URL}/api/jobs/{job_id}", headers=auth_headers, timeout=30)

    def test_create_job_default_ocalc(self, auth_headers):
        payload = {"name": "TEST_JT_DEFAULT"}
        r = requests.post(f"{BASE_URL}/api/jobs", json=payload, headers=auth_headers, timeout=30)
        assert r.status_code == 200
        job = r.json()
        try:
            assert job.get("job_type") == "OCALC"
        finally:
            requests.delete(f"{BASE_URL}/api/jobs/{job['id']}", headers=auth_headers, timeout=30)

    def test_update_job_type(self, auth_headers):
        # Create OCALC then update to SVCL/FS
        r = requests.post(f"{BASE_URL}/api/jobs", json={"name": "TEST_JT_UPDATE"}, headers=auth_headers, timeout=30)
        assert r.status_code == 200
        job_id = r.json()["id"]
        try:
            u = requests.put(f"{BASE_URL}/api/jobs/{job_id}", json={"job_type": "SVCL/FS"}, headers=auth_headers, timeout=30)
            assert u.status_code == 200, u.text
            assert u.json().get("job_type") == "SVCL/FS"
            # Verify via GET
            g = requests.get(f"{BASE_URL}/api/jobs/{job_id}", headers=auth_headers, timeout=30)
            assert g.json().get("job_type") == "SVCL/FS"
        finally:
            requests.delete(f"{BASE_URL}/api/jobs/{job_id}", headers=auth_headers, timeout=30)


# ============= ASE Data tests =============

class TestASEData:
    def test_save_and_get_ase_data(self, auth_headers, created_job_id):
        ase_payload = {
            "date_of_site_visit": "2026-01-15",
            "att_engineer": "John Engineer",
            "site_walked_by": "Jane Walker",
            "address": "123 Main St",
            "mpoe_address": "456 MPOE Ave",
            "occupancy_type": "Commercial",
            "number_of_floors": "5",
            "existing_inc": "Yes",
            "fiber_active_ports": "8",
            "nte_clli": "ATLNGAMA",
            "ef_offer": "Standard",
            "buried_pvc": "Yes",
            "buried_pvc_length": "100ft",
            "need_core_drill": "Yes",
            "core_drill_qty": "2",
            "pipelines_vicinity": "No",
            "additional_sow_notes": "All good",
        }
        r = requests.put(f"{BASE_URL}/api/jobs/{created_job_id}/ase-data", json=ase_payload, headers=auth_headers, timeout=30)
        assert r.status_code == 200, r.text

        g = requests.get(f"{BASE_URL}/api/jobs/{created_job_id}/ase-data", headers=auth_headers, timeout=30)
        assert g.status_code == 200
        data = g.json().get("ase_data", {})
        assert data.get("att_engineer") == "John Engineer"
        assert data.get("buried_pvc_length") == "100ft"
        assert data.get("core_drill_qty") == "2"

    def test_export_ase_excel(self, auth_headers, created_job_id):
        # First save some ASE data
        ase_payload = {
            "att_engineer": "Excel Engineer",
            "address": "999 Excel St",
            "buried_pvc": "Yes",
            "buried_pvc_length": "50ft",
        }
        requests.put(f"{BASE_URL}/api/jobs/{created_job_id}/ase-data", json=ase_payload, headers=auth_headers, timeout=30)

        r = requests.get(f"{BASE_URL}/api/jobs/{created_job_id}/export/ase", headers=auth_headers, timeout=60)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct or "excel" in ct.lower(), f"Unexpected content-type: {ct}"
        # Verify it's a valid xlsx
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Collect all cell values into a single string
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "ASE SITE WALK SURVEY" in text
        assert "EMPLOYEE INFO" in text
        assert "BUILDING CHARACTERISTICS" in text
        assert "FIBER ATTRIBUTES" in text
        assert "NTE - EQUIPMENT ATTRIBUTES" in text
        assert "EF REQUIREMENTS - DETAIL" in text
        assert "Excel Engineer" in text
        assert "999 Excel St" in text
        assert "50ft" in text

    def test_ase_data_unknown_job(self, auth_headers):
        r = requests.put(f"{BASE_URL}/api/jobs/non-existent-id/ase-data", json={"x": 1}, headers=auth_headers, timeout=30)
        assert r.status_code == 404
