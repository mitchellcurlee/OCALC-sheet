"""Backend tests for Daily Field Report (DFR) endpoints and xlsx export."""
import io
import os
import datetime as dt

import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "Admin@mastec.com"
ADMIN_PASS = "Mastecbyers12345"
USER_EMAIL = "mitchell.curlee@byers.com"
USER_PASS = "Dawson@2016"

SEED_JOB_ID = "e5dc0bae-96bb-4cef-8001-7baf4e41d80a"

EXPECTED_HEADERS = [
    "Date", "Job", "Market", "SOW", "Job Type", "Job Address", "Fielder at Site",
    "Travel to site (hrs)", "Work @job (hrs)", "Return to Home/next Site (hrs)",
    "Total miles driven", "Job footage (ft)", "# Poles", "Time on computer for Job",
    "Status", "Comments",
]


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=20)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="session")
def admin_token():
    return _login(ADMIN_EMAIL, ADMIN_PASS)


@pytest.fixture(scope="session")
def user_token():
    return _login(USER_EMAIL, USER_PASS)


def _h(t):
    return {"Authorization": f"Bearer {t}", "Content-Type": "application/json"}


# Use a test-only date far from today so we don't collide with smoke-test data
TEST_DATE = "2030-01-15"
TEST_DATE_2 = "2030-01-16"


@pytest.fixture(scope="module")
def admin_report(admin_token):
    """Create a DFR for admin on TEST_DATE; yield doc; cleanup."""
    body = {
        "date": TEST_DATE,
        "fielder": "Mitchell Curlee",
        "rows": [{
            "date": TEST_DATE, "job": SEED_JOB_ID, "market": "FTH", "sow": "OCALC",
            "job_type": "OCALC", "job_address": "24810 Majestic Cone Lane, Magnolia, tx 77355",
            "fielder": "Mitchell Curlee", "travel_hrs": 0.75, "work_hrs": 2, "return_hrs": 0.5,
            "miles": 34, "footage": 500, "poles": 3, "computer_hrs": 1,
            "status": "In Progress", "comments": "test row",
        }],
        "track": [],
    }
    r = requests.post(f"{API}/field-reports", json=body, headers=_h(admin_token), timeout=20)
    assert r.status_code == 201, f"create: {r.status_code} {r.text}"
    doc = r.json()
    yield doc
    # cleanup
    requests.delete(f"{API}/field-reports/{doc['id']}", headers=_h(admin_token), timeout=20)


class TestCrud:
    def test_create_201(self, admin_report):
        assert admin_report["date"] == TEST_DATE
        assert admin_report["fielder"] == "Mitchell Curlee"
        assert len(admin_report["rows"]) == 1
        assert admin_report["rows"][0]["market"] == "FTH"
        assert "id" in admin_report

    def test_second_post_same_date_is_idempotent(self, admin_token, admin_report):
        body = {"date": TEST_DATE, "fielder": "Other", "rows": [], "track": []}
        r = requests.post(f"{API}/field-reports", json=body, headers=_h(admin_token), timeout=20)
        # Spec: returns existing report (status may be 201 or 200)
        assert r.status_code in (200, 201)
        assert r.json()["id"] == admin_report["id"]
        # Not overwritten
        assert r.json()["fielder"] == "Mitchell Curlee"

    def test_list_own_filtered_by_date(self, admin_token, admin_report):
        r = requests.get(f"{API}/field-reports?date={TEST_DATE}", headers=_h(admin_token), timeout=20)
        assert r.status_code == 200
        ids = [d["id"] for d in r.json()]
        assert admin_report["id"] in ids

    def test_get_by_id(self, admin_token, admin_report):
        r = requests.get(f"{API}/field-reports/{admin_report['id']}", headers=_h(admin_token), timeout=20)
        assert r.status_code == 200
        assert r.json()["id"] == admin_report["id"]

    def test_put_replaces(self, admin_token, admin_report):
        body = {
            "date": TEST_DATE, "fielder": "Updated Fielder",
            "rows": [{"job": SEED_JOB_ID, "market": "FTD", "miles": 100}],
            "track": [{"lat": 30.1, "lng": -95.1, "t": 1, "row_id": "r1"}],
        }
        r = requests.put(f"{API}/field-reports/{admin_report['id']}", json=body, headers=_h(admin_token), timeout=20)
        assert r.status_code == 200
        updated = r.json()
        assert updated["fielder"] == "Updated Fielder"
        assert updated["rows"][0]["market"] == "FTD"
        # Verify persisted
        g = requests.get(f"{API}/field-reports/{admin_report['id']}", headers=_h(admin_token), timeout=20)
        assert g.json()["fielder"] == "Updated Fielder"

    def test_delete_and_404(self, admin_token):
        body = {"date": TEST_DATE_2, "fielder": "x", "rows": [], "track": []}
        r = requests.post(f"{API}/field-reports", json=body, headers=_h(admin_token), timeout=20)
        rid = r.json()["id"]
        d = requests.delete(f"{API}/field-reports/{rid}", headers=_h(admin_token), timeout=20)
        assert d.status_code == 200
        g = requests.get(f"{API}/field-reports/{rid}", headers=_h(admin_token), timeout=20)
        assert g.status_code == 404


class TestAccessControl:
    def test_other_user_can_get(self, user_token, admin_report):
        r = requests.get(f"{API}/field-reports/{admin_report['id']}", headers=_h(user_token), timeout=20)
        assert r.status_code == 200, f"regular user should GET admin's report: {r.status_code}"

    def test_other_user_cannot_put(self, user_token, admin_report):
        body = {"date": TEST_DATE, "fielder": "hacker", "rows": [], "track": []}
        r = requests.put(f"{API}/field-reports/{admin_report['id']}", json=body, headers=_h(user_token), timeout=20)
        assert r.status_code == 403

    def test_other_user_cannot_delete(self, user_token, admin_report):
        r = requests.delete(f"{API}/field-reports/{admin_report['id']}", headers=_h(user_token), timeout=20)
        assert r.status_code == 403

    def test_all_flag_lists_everyone(self, user_token, admin_report):
        r = requests.get(f"{API}/field-reports?all=1", headers=_h(user_token), timeout=20)
        assert r.status_code == 200
        ids = [d["id"] for d in r.json()]
        assert admin_report["id"] in ids, "all=1 must include other users' reports"

    def test_list_without_all_excludes_others(self, user_token, admin_report):
        r = requests.get(f"{API}/field-reports", headers=_h(user_token), timeout=20)
        assert r.status_code == 200
        ids = [d["id"] for d in r.json()]
        assert admin_report["id"] not in ids


class TestXlsxExport:
    def test_export_xlsx_headers_and_styling(self, admin_token, admin_report):
        # Re-create row (previous test overwrote via PUT)
        body = {
            "date": TEST_DATE, "fielder": "Mitchell Curlee",
            "rows": [{
                "date": TEST_DATE, "job": "A050ABC", "market": "FTH", "sow": "OCALC",
                "job_type": "OCALC", "job_address": "24810 Majestic Cone Lane",
                "fielder": "Mitchell Curlee", "travel_hrs": 0.75, "work_hrs": 2, "return_hrs": 0.5,
                "miles": 34, "footage": 500, "poles": 3, "computer_hrs": 1,
                "status": "In Progress", "comments": "test row",
            }],
            "track": [],
        }
        requests.put(f"{API}/field-reports/{admin_report['id']}", json=body, headers=_h(admin_token), timeout=20)

        r = requests.get(f"{API}/field-reports/{admin_report['id']}/export/xlsx", headers=_h(admin_token), timeout=30)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active

        # headers
        actual = [ws.cell(row=1, column=i + 1).value for i in range(len(EXPECTED_HEADERS))]
        assert actual == EXPECTED_HEADERS, f"headers mismatch: {actual}"

        # header styling: black fill, white bold font
        h = ws.cell(row=1, column=1)
        assert h.font.bold is True
        assert (h.font.color.rgb or "").upper().endswith("FFFFFF")
        assert (h.fill.fgColor.rgb or "").upper().endswith("000000")

        # alternating grey rows (row 2 is grey per build_dfr_workbook: ri%2==0 where ri starts at 0)
        row2 = ws.cell(row=2, column=1)
        row3 = ws.cell(row=3, column=1)
        c2 = (row2.fill.fgColor.rgb or "").upper()
        c3 = (row3.fill.fgColor.rgb or "").upper()
        assert c2.endswith("D9D9D9"), f"row2 fill expected D9D9D9, got {c2}"
        # row3 should NOT be grey
        assert not c3.endswith("D9D9D9"), f"row3 should not be grey, got {c3}"

        # data row values match (numbers as numbers)
        # Column indices: miles=11, footage=12, poles=13
        assert ws.cell(row=2, column=11).value == 34
        assert ws.cell(row=2, column=12).value == 500
        assert ws.cell(row=2, column=13).value == 3
        assert ws.cell(row=2, column=3).value == "FTH"  # market
