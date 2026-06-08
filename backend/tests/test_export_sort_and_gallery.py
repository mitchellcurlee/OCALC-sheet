"""Tests for new features:
  - Numerical sort of poles in GET /api/jobs/{job_id}/export/excel
  - Gallery enhancements (Download all / per-photo Download / open-in-new-tab)
  - Empty-state gallery ('No photos attached')
"""
import io
import os
import re
import uuid
from pathlib import Path

import pytest
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv(Path(__file__).resolve().parents[1] / ".env")

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
ADMIN_EMAIL = os.environ.get("MASTER_ADMIN_EMAIL")
ADMIN_PASSWORD = os.environ.get("MASTER_ADMIN_PASSWORD")
SEED_JOB_ID = os.environ.get("TEST_SEED_JOB_ID", "e5dc0bae-96bb-4cef-8001-7baf4e41d80a")

PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\xcf"
    b"\xc0\x00\x00\x00\x03\x00\x01\xa1\x06\x00>\x00\x00\x00\x00IEND\xaeB`\x82"
)


@pytest.fixture(scope="module")
def admin_headers():
    if not BASE_URL or not ADMIN_EMAIL or not ADMIN_PASSWORD:
        pytest.skip("env not configured")
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    if r.status_code != 200:
        pytest.skip(f"admin login failed: {r.text}")
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


class TestNumericExportSort:
    """Create poles 10, 2, 1 in seed job and assert export orders them 1, 2, 10."""
    created = []

    def test_create_unordered_poles(self, admin_headers):
        for num in ["10", "2", "1"]:
            payload = {
                "job_id": SEED_JOB_ID,
                "pole_number": num,
                "pole_size": "1-35",
                "condition": "Good",
            }
            r = requests.post(f"{BASE_URL}/api/poles", json=payload,
                              headers=admin_headers, timeout=30)
            assert r.status_code in (200, 201), r.text
            TestNumericExportSort.created.append(r.json()["id"])
        assert len(TestNumericExportSort.created) == 3

    def test_export_orders_numerically(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/jobs/{SEED_JOB_ID}/export/excel",
                         headers=admin_headers, timeout=60)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active

        # Collect all values in column B (which holds POLE # values).
        # POLE # cells in col A label; pole_number is at col B on the same row.
        pole_numbers_in_order = []
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "POLE #" and len(row) > 1:
                pole_numbers_in_order.append(str(row[1]) if row[1] is not None else "")

        # Filter the three we created (and any other numeric ones) to validate sort
        numeric = [p for p in pole_numbers_in_order if p.isdigit()]
        # Numeric ones must be ascending by integer value
        ints = [int(p) for p in numeric]
        assert ints == sorted(ints), f"numeric pole order wrong: {ints}"
        # 1,2,10 must appear in that relative order
        idx1 = numeric.index("1") if "1" in numeric else -1
        idx2 = numeric.index("2") if "2" in numeric else -1
        idx10 = numeric.index("10") if "10" in numeric else -1
        assert idx1 < idx2 < idx10, f"1,2,10 order broken: 1@{idx1} 2@{idx2} 10@{idx10}"

        # Non-numeric pole numbers (e.g., EXPORT-TEST*) should come AFTER numeric ones
        first_non_numeric_index = None
        last_numeric_index = None
        for i, p in enumerate(pole_numbers_in_order):
            if p and p[0].isdigit():
                last_numeric_index = i
            elif p and first_non_numeric_index is None:
                first_non_numeric_index = i
        if first_non_numeric_index is not None and last_numeric_index is not None:
            assert last_numeric_index < first_non_numeric_index, \
                "non-numeric pole appeared before a numeric one"

    def test_cleanup(self, admin_headers):
        for pid in TestNumericExportSort.created:
            requests.delete(f"{BASE_URL}/api/poles/{pid}",
                            headers=admin_headers, timeout=30)


class TestGalleryDownloadFeatures:
    pole_id_with_photos = None
    pole_id_no_photos = None

    def test_create_pole_with_photos(self, admin_headers):
        # upload one photo
        files = {"file": ("g.png", io.BytesIO(PNG_BYTES), "image/png")}
        up = requests.post(f"{BASE_URL}/api/upload-photo", files=files, timeout=60)
        assert up.status_code == 200
        url = up.json()["url"]

        payload = {
            "job_id": SEED_JOB_ID,
            "pole_number": f"TEST_GAL_{uuid.uuid4().hex[:5]}",
            "pole_size": "1-35", "condition": "Good",
            "ocalc_left_photo": url, "birthmark_photo": url, "mgnv_photo": url,
        }
        r = requests.post(f"{BASE_URL}/api/poles", json=payload,
                          headers=admin_headers, timeout=30)
        assert r.status_code in (200, 201)
        TestGalleryDownloadFeatures.pole_id_with_photos = r.json()["id"]

    def test_create_pole_no_photos(self, admin_headers):
        payload = {
            "job_id": SEED_JOB_ID,
            "pole_number": f"TEST_NOPHOTO_{uuid.uuid4().hex[:5]}",
            "pole_size": "1-35", "condition": "Good",
        }
        r = requests.post(f"{BASE_URL}/api/poles", json=payload,
                          headers=admin_headers, timeout=30)
        assert r.status_code in (200, 201)
        TestGalleryDownloadFeatures.pole_id_no_photos = r.json()["id"]

    def test_gallery_has_download_features(self):
        pid = TestGalleryDownloadFeatures.pole_id_with_photos
        assert pid
        r = requests.get(f"{BASE_URL}/api/poles/{pid}/gallery", timeout=30)
        assert r.status_code == 200
        html = r.text
        # Download all button
        assert 'class="dlall"' in html, "missing Download all button class"
        assert 'onclick="downloadAll()"' in html, "missing downloadAll handler"
        assert "Download all" in html
        # Per-photo Download anchor with download attribute
        dl_anchors = re.findall(r'<a\s+class="dl"[^>]*download="[^"]+"[^>]*>Download</a>', html)
        assert len(dl_anchors) >= 3, f"expected >=3 per-photo download links, got {len(dl_anchors)}"
        # Each photo wrapped in <a target="_blank">
        new_tab_anchors = re.findall(r'<a[^>]+target="_blank"[^>]*>\s*<img', html)
        assert len(new_tab_anchors) >= 3, f"expected >=3 open-in-new-tab links, got {len(new_tab_anchors)}"

    def test_gallery_empty_state(self):
        pid = TestGalleryDownloadFeatures.pole_id_no_photos
        assert pid
        r = requests.get(f"{BASE_URL}/api/poles/{pid}/gallery", timeout=30)
        assert r.status_code == 200
        html = r.text
        assert "No photos attached" in html
        assert 'class="dlall"' not in html  # button should not be rendered

    def test_cleanup(self, admin_headers):
        for pid in [TestGalleryDownloadFeatures.pole_id_with_photos,
                    TestGalleryDownloadFeatures.pole_id_no_photos]:
            if pid:
                requests.delete(f"{BASE_URL}/api/poles/{pid}",
                                headers=admin_headers, timeout=30)
