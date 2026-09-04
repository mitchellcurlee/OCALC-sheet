import requests, io, sys
from openpyxl import load_workbook

BASE = "https://utility-pole-map.preview.emergentagent.com/api"
tok = requests.post(f"{BASE}/auth/login", json={"email": "Admin@mastec.com", "password": "Mastecbyers12345"}).json()["access_token"]
H = {"Authorization": f"Bearer {tok}"}

TYPES = ["Handhole","Manhole","Brooks Enclosure","Flowerpot","Hoffman Box","Pedestal","Point of Intersection","FBT","NTE","FBT/NTE","PFP","SOC Start","SOC End","Riser","MPOE","EOP","ROW","DW Start","DW End","UG Path Start","UG Path End","Gas Line","Aerial Splice","Loop","Power","X-Messenger","Culvert","Drainage Start","Drainage End","Eyebolt","Pullbox","Weather Head","FDF","Gas Valve","Gas Meter","FST","Water Meter/Valve"]

job = requests.post(f"{BASE}/jobs", json={"site_id": "EXPORT-ALL-TEST", "name": "", "location": "1 Test St"}, headers=H).json()
jid = job["id"]
print("job", jid)

p = requests.post(f"{BASE}/poles", json={"job_id": jid, "pole_number": "T1", "latitude": 29.70, "longitude": -95.30}, headers=H)
print("pole", p.status_code)

lat = 29.70
for i, t in enumerate(TYPES):
    lat += 0.0002
    r = requests.post(f"{BASE}/underground", json={"job_id": jid, "type": t, "label": f"IT-{i}", "latitude": lat, "longitude": -95.30}, headers=H)
    if r.status_code not in (200, 201):
        print("CREATE FAIL", t, r.status_code, r.text[:150])

# Job Excel
r = requests.get(f"{BASE}/jobs/{jid}/export/excel", headers=H)
print("job excel", r.status_code)
ws = load_workbook(io.BytesIO(r.content)).active
rows = list(ws.iter_rows(values_only=True))
found, blank_abbr = {}, []
for row in rows:
    vals = [str(v) if v is not None else "" for v in row]
    if len(vals) > 4 and vals[4].startswith("IT-"):
        found[vals[1]] = vals[0]
        if not vals[0]:
            blank_abbr.append(vals[1])
missing = [t for t in TYPES if t not in found]
print("types in excel:", len(found), "/", len(TYPES))
print("MISSING:", missing)
print("BLANK ABBR:", blank_abbr)
# pole present?
pole_ok = any("T1" in [str(v) for v in row if v is not None] for row in rows)
print("pole T1 in excel:", pole_ok)

# ASE Excel
r2 = requests.get(f"{BASE}/jobs/{jid}/export/ase", headers=H)
print("ase excel", r2.status_code)
# SOW + Permission stamp
r3 = requests.post(f"{BASE}/jobs/{jid}/export/sow-docx", json={"text": "x"}, headers=H)
r4 = requests.get(f"{BASE}/jobs/{jid}/export/permission-stamp-docx", headers=H)
print("sow", r3.status_code, "ps", r4.status_code)

ok = not missing and not blank_abbr and pole_ok and r2.status_code == 200
requests.delete(f"{BASE}/jobs/{jid}", headers=H)
print("cleaned up job")
sys.exit(0 if ok else 1)
